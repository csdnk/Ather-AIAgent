#!/usr/bin/env python3
"""Generate P2 AetherEngine M0 review PowerPoint."""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

OUT = Path(__file__).resolve().parent.parent / "docs" / "P2_AetherEngine_M0_Presentation.pptx"

# Brand colors
GREEN = RGBColor(0x27, 0xAE, 0x60)
DARK = RGBColor(0x2C, 0x3E, 0x50)
BLUE = RGBColor(0x2E, 0x86, 0xC1)
ORANGE = RGBColor(0xE6, 0x7E, 0x22)
GRAY = RGBColor(0x7F, 0x8C, 0x8D)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)


def set_run(run, size=18, bold=False, color=DARK):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = "Microsoft YaHei"


def add_title_slide(prs, title, subtitle=""):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.shapes.add_shape(1, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = DARK
    bg.line.fill.background()

    accent = slide.shapes.add_shape(1, 0, Inches(4.8), prs.slide_width, Inches(0.08))
    accent.fill.solid()
    accent.fill.fore_color.rgb = GREEN
    accent.line.fill.background()

    tb = slide.shapes.add_textbox(Inches(0.8), Inches(1.6), Inches(11.5), Inches(1.5))
    p = tb.text_frame.paragraphs[0]
    r = p.add_run()
    r.text = title
    set_run(r, 36, True, WHITE)

    if subtitle:
        tb2 = slide.shapes.add_textbox(Inches(0.8), Inches(3.2), Inches(11.5), Inches(1.2))
        p2 = tb2.text_frame.paragraphs[0]
        r2 = p2.add_run()
        r2.text = subtitle
        set_run(r2, 20, False, GRAY)


def add_section_slide(prs, title):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bar = slide.shapes.add_shape(1, 0, 0, Inches(0.25), prs.slide_height)
    bar.fill.solid()
    bar.fill.fore_color.rgb = GREEN
    bar.line.fill.background()

    tb = slide.shapes.add_textbox(Inches(0.9), Inches(2.8), Inches(11), Inches(1.2))
    p = tb.text_frame.paragraphs[0]
    r = p.add_run()
    r.text = title
    set_run(r, 32, True, DARK)


def add_bullet_slide(prs, title, bullets, sub=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    # title bar
    bar = slide.shapes.add_shape(1, 0, 0, prs.slide_width, Inches(1.0))
    bar.fill.solid()
    bar.fill.fore_color.rgb = RGBColor(0xF8, 0xF9, 0xFA)
    bar.line.fill.background()

    tb = slide.shapes.add_textbox(Inches(0.6), Inches(0.22), Inches(12), Inches(0.7))
    p = tb.text_frame.paragraphs[0]
    r = p.add_run()
    r.text = title
    set_run(r, 24, True, GREEN)

    body = slide.shapes.add_textbox(Inches(0.7), Inches(1.3), Inches(12), Inches(5.5))
    tf = body.text_frame
    tf.word_wrap = True
    for i, item in enumerate(bullets):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.level = 0
        para.space_after = Pt(10)
        run = para.add_run()
        run.text = item
        set_run(run, 18, False, DARK)

    if sub:
        foot = slide.shapes.add_textbox(Inches(0.7), Inches(6.5), Inches(12), Inches(0.5))
        fp = foot.text_frame.paragraphs[0]
        fr = fp.add_run()
        fr.text = sub
        set_run(fr, 12, False, GRAY)


def add_table_slide(prs, title, headers, rows):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    tb = slide.shapes.add_textbox(Inches(0.6), Inches(0.3), Inches(12), Inches(0.7))
    p = tb.text_frame.paragraphs[0]
    r = p.add_run()
    r.text = title
    set_run(r, 24, True, GREEN)

    cols, rs = len(headers), len(rows) + 1
    table = slide.shapes.add_table(rs, cols, Inches(0.5), Inches(1.2), Inches(12.3), Inches(0.45 * rs)).table

    for c, h in enumerate(headers):
        cell = table.cell(0, c)
        cell.text = h
        cell.fill.solid()
        cell.fill.fore_color.rgb = GREEN
        for para in cell.text_frame.paragraphs:
            para.alignment = PP_ALIGN.CENTER
            for run in para.runs:
                set_run(run, 14, True, WHITE)

    for ri, row in enumerate(rows, start=1):
        for ci, val in enumerate(row):
            cell = table.cell(ri, ci)
            cell.text = val
            for para in cell.text_frame.paragraphs:
                para.alignment = PP_ALIGN.CENTER
                for run in para.runs:
                    set_run(run, 13, False, DARK)


def add_closing_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.shapes.add_shape(1, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = GREEN
    bg.line.fill.background()

    tb = slide.shapes.add_textbox(Inches(1), Inches(2.5), Inches(11), Inches(2))
    tf = tb.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = "谢谢聆听"
    set_run(r, 40, True, WHITE)

    p2 = tf.add_paragraph()
    p2.alignment = PP_ALIGN.CENTER
    r2 = p2.add_run()
    r2.text = "P2 AetherEngine · M0 阶段评审 · 南京信息工程大学"
    set_run(r2, 18, False, WHITE)


def main():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    add_title_slide(
        prs,
        "P2 AetherEngine 六态存储引擎",
        "M0 阶段评审汇报 | 2026-07-31 | 南京信息工程大学",
    )

    add_bullet_slide(
        prs,
        "汇报提纲",
        [
            "1. 项目背景与合同范围",
            "2. P2 总体架构",
            "3. E1 / E2 / E3 三引擎方案",
            "4. M0 接口冻结（API + 内省 + 迁移回调）",
            "5. 第一阶段交付物清单",
            "6. 自测结果与下一步计划",
        ],
    )

    add_section_slide(prs, "01  项目背景与合同范围")

    add_bullet_slide(
        prs,
        "项目定位",
        [
            "AetherStore：AI 原生超融合统一存储平台",
            "P2 AetherEngine：存储引擎层，位于 P1 底座之上、P3/P4 之下",
            "本阶段聚焦三态：E1 向量 · E2 对象 · E3 图",
            "E4 文件 / E5 块 / E6 时序：M0 仅调研，不实现",
        ],
        "合同：捷云盛 × 南信大 | 横向合同 2026-06-05 ~ 2027-06-30",
    )

    add_table_slide(
        prs,
        "合同里程碑",
        ["里程碑", "时间", "目标"],
        [
            ["M0", "2026-07-31", "方案 + 接口冻结"],
            ["M-MVP", "2026-11-30", "E1/E2 可生产 + P1 接入"],
            ["M1", "2027-03-31", "E3 融合 + DiskANN + EC"],
            ["三态 GA", "2027-06-30", "生产可用 + Bench + 论文专利"],
        ],
    )

    add_section_slide(prs, "02  P2 总体架构")

    add_bullet_slide(
        prs,
        "架构分层",
        [
            "P4 接入层：S3 / gRPC / REST / PG Wire 协议网关",
            "P3 语义层：Embedding 下推 · 记忆管理 · 智能分层调度",
            "P2 引擎层：统一 API + Segment 生命周期 + WAL + 错误码",
            "P1 底座：Storage Block · WAL · Raft · Tiering Hook",
        ],
    )

    add_bullet_slide(
        prs,
        "共享底座（三引擎统一）",
        [
            "Segment 生命周期：Growing → Sealing → Sealed → Frozen → Archived",
            "WAL & Recovery：magic + length + CRC32，崩溃可 replay",
            "P2Err_* 统一错误码，线稳定对外暴露",
            "SegmentControl：Freeze / Unfreeze / OnMigrateComplete / Failed",
            "OTel：ae.<engine>.<op> span，M-MVP 接入 P1 总线",
        ],
    )

    add_section_slide(prs, "03  E1 / E2 / E3 三引擎")

    add_table_slide(
        prs,
        "三引擎能力概览",
        ["引擎", "数据模型", "核心能力", "状态"],
        [
            ["E1 向量", "Collection/Vector", "HNSW + IVF-PQ + Segment + 删除", "已实现"],
            ["E2 对象", "Bucket/Object", "S3 子集 + Range + 分页 + Multipart", "已实现"],
            ["E3 图", "Property Graph", "遍历 + 向量投影 + 融合算子", "已实现"],
        ],
    )

    add_bullet_slide(
        prs,
        "E1 向量引擎亮点",
        [
            "Segment 分段存储，Growing 段 Flat 回退，Sealed 段建索引",
            "HNSW 纯 Rust ANN 主索引",
            "IVF-PQ 第二索引 + 自适应参数（nlist≈√n，recall≥92% 门禁通过）",
            "Delete + tombstone + WAL 恢复",
        ],
    )

    add_bullet_slide(
        prs,
        "E2 / E3 亮点",
        [
            "E2：ObjectBackend 抽象，LocalFS 生产可用，SeaweedFS 预留",
            "E2：MD5 ETag + blake3，object key 安全哈希映射",
            "E3：VectorAnchoredSubgraph 单读路径（锚点 + k-hop + GraphFilter）",
            "E3：truncated 标记 + 写栅栏，Frozen 期只读可服务",
        ],
    )

    add_section_slide(prs, "04  M0 接口冻结")

    add_bullet_slide(
        prs,
        "proto/aether_engine.proto — FROZEN v0.1",
        [
            "① 三态统一查询 API：Vector / Object / Graph / Fusion Service",
            "② Segment 内省：SegmentStats + ListSegments（IF-04 → P3）",
            "③ Tier Migrate Callback：Freeze / Unfreeze / OnMigrate*（IF-06）",
            "冻结规约：字段号与 RPC 签名稳定，仅允许向后兼容新增",
        ],
        "Rust 镜像：crates/ae-proto | 实现：ae-kernel::SegmentControl",
    )

    add_section_slide(prs, "05  第一阶段交付物")

    add_table_slide(
        prs,
        "M0 文档交付物（12 项）",
        ["序号", "交付物", "状态"],
        [
            ["D-01", "需求分析文档", "✅"],
            ["D-02", "总体架构设计", "✅"],
            ["D-03", "阶段方案（E1/E2/E3）", "✅"],
            ["D-04", "技术路线与实施计划", "✅"],
            ["D-05", "E4/E5/E6 调研报告", "✅"],
            ["D-06", "P1/P2/P3 接口依赖表", "✅"],
            ["D-07", "API 接口文档 v0.1", "✅"],
            ["D-08~12", "proto + 代码 + 自测报告", "✅"],
        ],
    )

    add_section_slide(prs, "06  自测与展望")

    add_bullet_slide(
        prs,
        "自测结果（2026-06-25）",
        [
            "cargo fmt / clippy：零 warning",
            "cargo test --workspace：31 passed, 0 failed",
            "ae-server demo：向量写入 + 图融合单路径跑通",
            "IVF-PQ recall@10 ≥ 0.92 合同门禁通过",
        ],
    )

    add_bullet_slide(
        prs,
        "下一步（M-MVP · 2026-08 ~ 11）",
        [
            "真实 P1 Block/WAL 接入（替换 MockP1Client）",
            "E1/E2 生产化性能压测（P99 / 吞吐 / 5 千万召回）",
            "P4 网关联调（S3 / gRPC）",
            "OTel 导出接入 P1 监控总线",
            "CI 流水线（fmt + clippy + test）",
        ],
    )

    add_bullet_slide(
        prs,
        "提请甲方确认",
        [
            "Tier Migrate 物理迁移由 P1 还是 P3 执行？",
            "P1 Block/WAL SDK v0.1 何时冻结？",
            "P3-B1 Embedding 写入字段是否与 VectorService 对齐？",
            "GA 前 E4/E5/E6 是否仅需调研报告？",
        ],
    )

    add_closing_slide(prs)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(OUT))
    print(f"Generated: {OUT}")


if __name__ == "__main__":
    main()
